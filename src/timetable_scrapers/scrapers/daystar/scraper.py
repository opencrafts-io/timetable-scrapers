from datetime import datetime
from typing import List
from openpyxl import load_workbook

from ...base.scraper import BaseTimetableScraper
from ...registry import ScraperRegistry
from ...schemas import CourseEntry
from ...utils.time_parser import parse_exam_datetime, calculate_duration


@ScraperRegistry.register("daystar")
class SchoolExamScraper(BaseTimetableScraper):
    """
    Scraper for Daystar University Exams.
    Matches the minimal Professor API contract.
    Handles multiple worksheets, each with room information in the first column.
    """

    @property
    def institution_name(self) -> str:
        return "Daystar University"

    @property
    def timezone(self) -> str:
        return "EAT"  # East Africa Time

    def extract(self, file) -> List[CourseEntry]:
        """Extract exam timetable data from school Excel file."""
        wb_obj = load_workbook(filename=file)
        work_sheets = wb_obj.sheetnames

        courses = []

        days_of_the_week = [
            "MONDAY",
            "TUESDAY",
            "WEDNESDAY",
            "THURSDAY",
            "FRIDAY",
            "SATURDAY",
        ]

        for sheet in work_sheets:
            work_sheet = wb_obj[sheet]

            # Reset per sheet: room-by-row-index is only valid within the
            # sheet it was read from. Sharing this across sheets let a
            # room name from one campus's sheet leak into another
            # campus's course entries whenever row-index ranges overlapped.
            rooms = {}

            for column_one in work_sheet.iter_cols(values_only=True):
                for i, room in enumerate(column_one):
                    if room is None or room == "ROOM":
                        continue
                    rooms[f"{i}"] = room
                break

            data_columns = list(work_sheet.iter_cols(values_only=True))[1:]

            def is_day_value(value) -> bool:
                return isinstance(value, datetime) or (
                    isinstance(value, str)
                    and any(d in value.upper() for d in days_of_the_week)
                )

            def to_day(value) -> str:
                return value.strftime("%Y-%m-%d") if isinstance(value, datetime) else value

            # The date header is only written once, in the leftmost ("anchor")
            # column of each day-block; the other time-slot columns belonging
            # to that block share the date but have no date cell of their
            # own. Track each column's own day-by-row, then have columns
            # without any date cells inherit the nearest anchor column's
            # (to their left) day-by-row, so the date lines up by row
            # position rather than leaking the last date seen while
            # scanning a previous column top-to-bottom.
            anchor_day_by_row = []
            course_time_range = ""
            course_code = ""

            for column in data_columns:
                column_day_by_row = []
                running_day = ""
                for value in column:
                    if is_day_value(value):
                        running_day = to_day(value)
                    column_day_by_row.append(running_day)

                if any(column_day_by_row):
                    anchor_day_by_row = column_day_by_row
                day_by_row = anchor_day_by_row or column_day_by_row

                for idx, value in enumerate(column):
                    if value is None:
                        continue

                    if value == "CHAPEL":
                        continue

                    if is_day_value(value):
                        pass
                    elif (
                        isinstance(value, str)
                        and len(value) > 0
                        and value[0].isdigit()
                    ):
                        course_time_range = value.strip()
                    elif isinstance(value, str):
                        course_code = value
                        day = day_by_row[idx]

                        # Parse time range and create entry
                        start_iso = ""
                        end_iso = ""
                        if "-" in course_time_range:
                            time_parts = course_time_range.split("-", 1)
                            start_iso = parse_exam_datetime(day, time_parts[0].strip(), self.timezone)
                            end_iso = parse_exam_datetime(day, time_parts[1].strip(), self.timezone)

                        if not start_iso or not end_iso:
                            continue

                        hrs = calculate_duration(start_iso, end_iso)

                        courses.append(
                            CourseEntry(
                                course_code=course_code,
                                start_time=start_iso,
                                end_time=end_iso,
                                venue=str(rooms.get(f"{idx}", "")).strip() or "TBA",
                                hrs=str(hrs),
                                raw_data={
                                    "original_day": day,
                                    "original_time": course_time_range,
                                }
                            )
                        )

        return self.normalize_output(courses)
