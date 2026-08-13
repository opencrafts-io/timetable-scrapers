from io import BytesIO

from openpyxl import Workbook

from timetable_scrapers.scrapers.daystar.scraper import SchoolExamScraper


def _build_workbook():
    """
    Mimics the real Daystar layout: the date header is only written once,
    in the leftmost ("anchor") column of each day-block, and is *not*
    repeated in the other time-slot columns that belong to the same block.

    Column A: rooms (one per row)
    Column B: anchor column - carries its own date header before each block
    Column C: follower column - shares B's dates but has no date cells itself
    """
    wb = Workbook()
    ws = wb.active

    rows = [
        ["ROOM1", "MONDAY 17/08/2026", None],
        [None, "9:00AM-11:00AM", "11:30AM-1:30PM"],
        ["ROOM2", "AAA111A", "CCC333A"],
        [None, "MONDAY 24/08/2026", None],
        [None, "9:00AM-11:00AM", "11:30AM-1:30PM"],
        ["ROOM3", "BBB222A", "DDD444A"],
    ]
    for row in rows:
        ws.append(row)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_follower_column_inherits_anchor_column_date_per_row():
    scraper = SchoolExamScraper()
    entries = scraper.extract(_build_workbook())

    by_code = {e.course_code: e for e in entries}

    assert by_code["AAA111A"].start_time.startswith("2026-08-17")
    assert by_code["CCC333A"].start_time.startswith("2026-08-17")
    assert by_code["BBB222A"].start_time.startswith("2026-08-24")
    assert by_code["DDD444A"].start_time.startswith("2026-08-24")


def _build_multi_sheet_workbook():
    """
    Two sheets (different campuses). Sheet A has a room name at row-index 2.
    Sheet B has a course at the same row-index 2 but its own room column is
    blank there - the room-by-row-index map must not leak from Sheet A.
    """
    wb = Workbook()

    ws_a = wb.active
    ws_a.title = "A"
    for row in [
        ["AROOM0", "MONDAY 17/08/2026"],
        [None, "9:00AM-11:00AM"],
        ["AROOM2", "AAA111A"],
    ]:
        ws_a.append(row)

    ws_b = wb.create_sheet("B")
    for row in [
        [None, "TUESDAY 18/08/2026"],
        [None, "9:00AM-11:00AM"],
        [None, "XXX999A"],
    ]:
        ws_b.append(row)

    buf = BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf


def test_room_lookup_does_not_leak_across_sheets():
    scraper = SchoolExamScraper()
    entries = scraper.extract(_build_multi_sheet_workbook())

    by_code = {e.course_code: e for e in entries}

    assert by_code["AAA111A"].venue == "AROOM2"
    assert by_code["XXX999A"].venue == "TBA"
